# -*- coding: utf-8 -*-
"""สร้างคู่มือพัฒนาระบบ SCiGENIE แบบ SDLC สรุป ใช้งานได้จริง สำหรับผู้พัฒนา"""
from __future__ import annotations

import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = os.path.join("HR", "SCiGENIE-คู่มือพัฒนาระบบ-SDLC.docx")
FLOW = os.path.join("HR", "SCiGENIE-system-flow.png")
ER = os.path.join("HR", "SCiGENIE-relational-db-design.png")
FONT = "Cordia New"  # มีบน Windows มาตรฐาน; ถ้ามี TH Sarabun New สามารถเปลี่ยนได้


def set_run_font(run, size=16, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT
    r = run._element.get_or_add_rPr()
    rFonts = r.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = color


def add_para(doc, text, *, size=16, bold=False, center=False, space_after=6, first_line=None):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line is not None:
        pf.first_line_indent = Cm(first_line)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_h(doc, text, level=1):
    # Heading style + Thai font
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=18 if level == 1 else 16, bold=True)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_bullets(doc, items, size=16):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(2)


def add_numbered(doc, items, size=16):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=14, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=14)
    doc.add_paragraph()


def add_caption(doc, text):
    add_para(doc, text, size=14, center=True, space_after=10)


def add_image(doc, path, width_cm=15.5):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Cm(width_cm))
        return True
    return False


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2)

    # ===== Cover =====
    add_para(doc, "คู่มือการพัฒนาระบบตามวงจรชีวิตซอฟต์แวร์ (SDLC)", size=22, bold=True, center=True, space_after=4)
    add_para(doc, "ระบบตอบคำถามอัจฉริยะ SCiGENIE", size=20, bold=True, center=True, space_after=4)
    add_para(doc, "Science Intelligent Enquiry", size=16, center=True, space_after=12)
    add_para(doc, "ฉบับสำหรับผู้พัฒนาระบบ — สรุป เข้าใจง่าย ใช้งานได้จริง", size=16, center=True, space_after=18)
    add_para(doc, "คณะวิทยาศาสตร์ มหาวิทยาลัยขอนแก่น", size=16, center=True, space_after=2)
    add_para(doc, "Production: https://sci.kku.ac.th/app/SCiGENIE/", size=14, center=True, space_after=2)
    add_para(doc, "Admin: https://sci.kku.ac.th/app/SCiGENIE/admin/", size=14, center=True, space_after=18)
    add_para(doc, "จัดทำโดย นางจรูญลักษณ์ ตรงดี  นักวิชาการคอมพิวเตอร์ชำนาญการ", size=14, center=True, space_after=2)
    add_para(doc, "กลุ่มภาระงานด้านเทคโนโลยีสารสนเทศ", size=14, center=True, space_after=24)

    add_para(
        doc,
        "เอกสารนี้สรุปกระบวนการพัฒนาระบบตาม SDLC ตั้งแต่วิเคราะห์ความต้องการ ออกแบบ พัฒนา ทดสอบ ติดตั้ง จนถึงบำรุงรักษา "
        "เน้นข้อมูลที่ผู้พัฒนานำไปใช้จริง และตัดรายละเอียดที่ไม่จำเป็นต่อการพัฒนาออก",
        size=15,
        space_after=12,
    )

    # ===== TOC =====
    add_h(doc, "สารบัญ", 1)
    add_bullets(
        doc,
        [
            "บทที่ 1  บทนำและภาพรวมระบบ",
            "บทที่ 2  วิเคราะห์ความต้องการ (Analysis)",
            "บทที่ 3  ออกแบบระบบ (Design) และพจนานุกรมข้อมูล (Data Dictionary)",
            "บทที่ 4  พัฒนาและเชื่อมต่อระบบ (Implementation)",
            "บทที่ 5  ทดสอบระบบ (Testing)",
            "บทที่ 6  ติดตั้งขึ้นเครื่องแม่ข่าย (Deployment)",
            "บทที่ 7  บำรุงรักษาและปรับปรุงต่อเนื่อง (Maintenance)",
            "ภาคผนวก  ตารางอ้างอิงด่วน",
        ],
    )

    # ===== Ch1 =====
    add_h(doc, "บทที่ 1  บทนำและภาพรวมระบบ", 1)
    add_h(doc, "1.1 จุดประสงค์ของคู่มือ", 2)
    add_para(
        doc,
        "คู่มือนี้ใช้เป็นแนวทางสำหรับนักพัฒนาระบบในการติดตั้ง พัฒนาต่อ ทดสอบ และดูแล SCiGENIE "
        "โดยยึดหลักว่า “ตอบจากคลังความรู้ของคณะเท่านั้น ไม่แต่งคำตอบเอง”",
    )
    add_h(doc, "1.2 ระบบคืออะไร", 2)
    add_para(
        doc,
        "SCiGENIE (Science Intelligent Enquiry) เป็นระบบตอบคำถามอัตโนมัติของคณะวิทยาศาสตร์ "
        "ผู้ใช้ถามผ่านเว็บไซต์ Facebook Messenger หรือ LINE ระบบค้นคำตอบจากคลัง FAQ เดียวกัน "
        "ให้คะแนนความใกล้เคียง แล้วให้ AI ช่วยตรวจก่อนตอบ หากไม่มั่นใจจะไม่เดา แต่แนะนำเลือกหมวดหรือติดต่อเจ้าหน้าที่",
    )
    add_h(doc, "1.3 เป้าหมายหลัก", 2)
    add_bullets(
        doc,
        [
            "ให้บริการข้อมูลพื้นฐานได้ตลอด 24 ชั่วโมง ผ่าน 3 ช่องทาง",
            "ลดภาระตอบคำถามซ้ำของเจ้าหน้าที่",
            "ให้เจ้าของหมวดปรับปรุงคลังคำถามเองได้ตามสิทธิ์",
            "เก็บสถิติคำถามจริงเพื่อเติม FAQ ที่ยังขาด",
        ],
    )
    add_h(doc, "1.4 วงจร SDLC ที่ใช้", 2)
    add_para(
        doc,
        "โครงการใช้แนวทาง SDLC แบบวนซ้ำได้ (Iterative) ตามลำดับ: "
        "วิเคราะห์ → ออกแบบ → พัฒนา → ทดสอบ → ติดตั้ง → บำรุงรักษา "
        "เมื่อมีคำถามใหม่หรือเปลี่ยนเกณฑ์คะแนน ให้วนกลับไปทดสอบและปรับปรุงคลังความรู้",
    )

    # ===== Ch2 =====
    add_h(doc, "บทที่ 2  วิเคราะห์ความต้องการ (Analysis)", 1)
    add_h(doc, "2.1 ปัญหาที่ต้องแก้", 2)
    add_bullets(
        doc,
        [
            "คำถามพื้นฐานซ้ำ ๆ ใช้เวลเจ้าหน้าที่มาก",
            "ผู้ใช้ถามนอกเวลาราชการแล้วต้องรอ",
            "ตอบหลายช่องทางโดยหลายคน อาจได้ข้อมูลไม่ตรงกัน หากไม่มีคลังกลาง",
        ],
    )
    add_h(doc, "2.2 ความต้องการเชิงหน้าที่ (Functional)", 2)
    add_table(
        doc,
        ["กลุ่ม", "ต้องทำได้"],
        [
            ["ผู้ใช้ทั่วไป", "เลือกหมวด / พิมพ์คำถามไทย / ได้คำตอบ หรือคำแนะนำติดต่อเจ้าหน้าที่"],
            ["ช่องทาง", "เว็บ, Facebook Messenger, LINE ใช้กลไกค้นหาชุดเดียวกัน"],
            ["ผู้ดูแลหมวด", "เพิ่ม–แก้ไข–ลบ FAQ เฉพาะหมวดที่ได้รับสิทธิ์"],
            ["ผู้ดูแลระบบ (super)", "จัดการผู้ใช้ สิทธิ์หมวด ดูสถิติทุกหมวด"],
            ["ระบบ", "ให้คะแนน → AI ตรวจ → ตอบหรือไม่ตอบตามเกณฑ์"],
        ],
    )
    add_h(doc, "2.3 ความต้องการเชิงคุณภาพ (Non-functional)", 2)
    add_bullets(
        doc,
        [
            "ถูกต้อง: ไม่ตอบนอกคลังความรู้",
            "ปลอดภัย: ตรวจลายเซ็น webhook, HTTPS, สิทธิ์น้อยที่สุด",
            "พร้อมใช้: ติดตั้งบนเซิร์ฟเวอร์คณะภายใต้โดเมน มข.",
            "ดูแลต่อได้: แยกค่าความลับออกจากโค้ด (.env)",
        ],
    )
    add_h(doc, "2.4 สิ่งที่อยู่นอกขอบเขต", 2)
    add_bullets(
        doc,
        [
            "ไม่ใช่ระบบทำธุรกรรม (ลงทะเบียน / ชำระเงิน / อนุมัติคำร้อง)",
            "ไม่แทนดุลยพินิจเจ้าหน้าที่ในกรณีพิเศษ",
            "ไม่แต่งระเบียบใหม่ด้วย AI",
        ],
    )

    # ===== Ch3 =====
    add_h(doc, "บทที่ 3  ออกแบบระบบ (Design)", 1)
    add_h(doc, "3.1 สถาปัตยกรรมโดยสรุป", 2)
    add_para(
        doc,
        "สถาปัตยกรรมแบบหลายช่องทาง ใช้คลังความรู้ชุดเดียว แบ่งชั้นงานดังนี้: "
        "ช่องทางรับข้อความ → API/แอป → กลไกค้นหา+AI → ฐานข้อมูล FAQ → บันทึกสถิติ",
    )
    add_bullets(
        doc,
        [
            "ชั้นช่องทาง: widget เว็บ, LINE webhook, Facebook webhook",
            "ชั้นแอป: PHP (Matcher, LineBot, MessengerBot, Auth)",
            "ชั้นความรู้: MySQL (categories, faqs, faq_embeddings, admin_users)",
            "ชั้น AI: ตรวจความเหมาะสมของคำตอบจากคลัง (ไม่แต่งคำตอบ)",
            "ชั้นปฏิบัติการ: AskLog / session หมวดของผู้ใช้",
        ],
    )
    add_h(doc, "3.2 แผนภาพการทำงาน (Flow)", 2)
    add_para(doc, "ภาพด้านล่างแสดงลำดับตั้งแต่มือผู้ใช้ถาม จนถึงผลลัพธ์ และเกณฑ์คะแนนตอบ/ไม่ตอบ")
    if add_image(doc, FLOW, 15.8):
        add_caption(doc, "ภาพที่ 1  แผนภาพการทำงานของระบบ SCiGENIE")
    add_h(doc, "3.3 เกณฑ์คะแนน (ต้องคงไว้เมื่อแก้โค้ด)", 2)
    add_table(
        doc,
        ["คะแนน", "ความหมาย", "ผลที่ควรเกิด"],
        [
            ["ต่ำกว่า 12", "ไม่คล้ายพอ", "ไม่ตอบจากคลัง → ขอเลือกหมวด หรือแนะนำเจ้าหน้าที่"],
            ["12 – 21", "มีเค้าแต่ยังไม่มั่นใจ", "ตอบได้เมื่อ AI ยืนยัน หรือมีสัญญาณชัด"],
            ["ตั้งแต่ 22", "มั่นใจปานกลาง", "พร้อมตอบ"],
            ["ตั้งแต่ 40", "มั่นใจสูง", "ตอบได้ชัดเจน"],
        ],
    )
    add_para(
        doc,
        "หมายเหตุสำหรับผู้พัฒนา: ค่าในโค้ดอ้างอิง MIN_SCORE = 12 และระดับเชื่อถือที่ 22 / 40 "
        "ห้ามคัดลอกเกณฑ์คะแนนไปแยกใน LINE/Facebook ให้เรียก Matcher ตัวเดียวกันเสมอ",
        size=15,
    )
    add_h(doc, "3.4 ออกแบบฐานข้อมูล", 2)
    add_para(
        doc,
        "ฐานข้อมูลหลักชื่อ scigeniedb แยกจากฐานบุคลากร (eoffice) "
        "แอปอ่านข้อมูลบุคลากรเฉพาะตอนตรวจสิทธิ์เข้าแอดมิน",
    )
    add_table(
        doc,
        ["ตาราง", "หน้าที่"],
        [
            ["categories", "หมวดคำถาม"],
            ["faqs", "คำถาม คำสำคัญ คำตอบ ผูกหมวด"],
            ["faq_embeddings", "เวกเตอร์ความหมายสำหรับค้นเชิง semantic (1:1 กับ faqs)"],
            ["admin_users", "ผู้ดูแล (role = super | manager)"],
            ["admin_user_categories", "สิทธิ์หมวดของผู้จัดการ (M:N)"],
        ],
    )
    if add_image(doc, ER, 15.8):
        add_caption(doc, "ภาพที่ 2  แผนภาพออกแบบฐานข้อมูลเชิงสัมพันธ์ (ER)")

    add_h(doc, "3.5 พจนานุกรมข้อมูล (Data Dictionary)", 2)
    add_para(
        doc,
        "ฐานข้อมูล scigeniedb ใช้ชุดอักขระ utf8mb4 "
        "ตารางด้านล่างสรุปฟิลด์ที่ผู้พัฒนาต้องรู้จักเมื่อออกแบบ API หรือแก้โค้ด",
    )

    add_para(doc, "3.5.1 ตาราง categories — หมวดคำถาม", size=15, bold=True, space_after=4)
    add_table(
        doc,
        ["ฟิลด์", "ชนิดข้อมูล", "คีย์/ข้อกำหนด", "คำอธิบาย"],
        [
            ["id", "INT UNSIGNED", "PK, AI", "รหัสหมวด"],
            ["name", "VARCHAR(191)", "UNIQUE, NOT NULL", "ชื่อหมวดเต็ม ใช้แสดงผลและผูกสิทธิ์"],
            ["short_name", "VARCHAR(80)", "DEFAULT ''", "ชื่อสั้นสำหรับป้าย/ชิป"],
            ["icon", "VARCHAR(40)", "DEFAULT folder", "ชื่อไอคอนในวิดเจ็ต"],
            ["hint", "VARCHAR(191)", "DEFAULT ''", "คำใบ้หมวดสำหรับผู้ใช้"],
            ["sort_order", "INT", "DEFAULT 0", "ลำดับการแสดง"],
            ["created_at", "TIMESTAMP", "DEFAULT NOW", "วันเวลาสร้าง"],
            ["updated_at", "TIMESTAMP", "ON UPDATE", "วันเวลาแก้ไขล่าสุด"],
        ],
    )

    add_para(doc, "3.5.2 ตาราง faqs — คำถาม–คำตอบในคลัง", size=15, bold=True, space_after=4)
    add_table(
        doc,
        ["ฟิลด์", "ชนิดข้อมูล", "คีย์/ข้อกำหนด", "คำอธิบาย"],
        [
            ["id", "INT UNSIGNED", "PK, AI", "รหัส FAQ"],
            ["category_id", "INT UNSIGNED", "FK → categories.id", "หมวดเจ้าของรายการ (ON DELETE RESTRICT)"],
            ["question", "TEXT", "NOT NULL", "ข้อความคำถามมาตรฐาน"],
            ["keyword", "TEXT", "NOT NULL", "คำสำคัญหลัก (ภาษาทางการ)"],
            ["keyword_back", "TEXT", "NOT NULL", "คำสำคัญเสริม (ภาษาพูด/คำย่อ)"],
            ["answer", "TEXT", "NOT NULL", "คำตอบที่แสดงให้ผู้ใช้"],
            ["created_at", "TIMESTAMP", "DEFAULT NOW", "วันเวลาสร้าง"],
            ["updated_at", "TIMESTAMP", "ON UPDATE", "วันเวลาแก้ไขล่าสุด"],
        ],
    )

    add_para(doc, "3.5.3 ตาราง faq_embeddings — ดัชนีความหมาย", size=15, bold=True, space_after=4)
    add_table(
        doc,
        ["ฟิลด์", "ชนิดข้อมูล", "คีย์/ข้อกำหนด", "คำอธิบาย"],
        [
            ["faq_id", "INT UNSIGNED", "PK, FK → faqs.id", "ผูก 1:1 กับ FAQ (ON DELETE CASCADE)"],
            ["model", "VARCHAR(120)", "DEFAULT ''", "ชื่อรุ่นเวกเตอร์ เช่น local-hash-v2"],
            ["vector_json", "MEDIUMTEXT", "NOT NULL", "เวกเตอร์ความหมายเก็บเป็น JSON"],
            ["content_hash", "VARCHAR(64)", "DEFAULT ''", "แฮชเนื้อหา ใช้รู้ว่าต้อง rebuild หรือไม่"],
            ["updated_at", "TIMESTAMP", "ON UPDATE", "วันเวลาอัปเดตเวกเตอร์"],
        ],
    )

    add_para(doc, "3.5.4 ตาราง admin_users — ผู้ดูแลระบบ", size=15, bold=True, space_after=4)
    add_table(
        doc,
        ["ฟิลด์", "ชนิดข้อมูล", "คีย์/ข้อกำหนด", "คำอธิบาย"],
        [
            ["id", "INT UNSIGNED", "PK, AI", "รหัสผู้ดูแล"],
            ["email", "VARCHAR(191)", "UNIQUE, NOT NULL", "อีเมล Google ที่ใช้ล็อกอิน"],
            ["eoffice_username", "VARCHAR(64)", "DEFAULT ''", "username ในระบบบุคลากร (เชื่อมเชิงตรรกะ)"],
            ["name", "VARCHAR(191)", "DEFAULT ''", "ชื่อที่แสดง"],
            ["role", "ENUM(super,manager)", "DEFAULT manager", "super = ทุกหมวด / manager = ตามสิทธิ์"],
            ["is_active", "TINYINT(1)", "DEFAULT 1", "1 = ใช้งานได้, 0 = ระงับ"],
            ["last_login_at", "DATETIME", "NULL", "เข้าสู่ระบบล่าสุด"],
            ["created_at", "TIMESTAMP", "DEFAULT NOW", "วันเวลาสร้าง"],
            ["updated_at", "TIMESTAMP", "ON UPDATE", "วันเวลาแก้ไขล่าสุด"],
        ],
    )

    add_para(doc, "3.5.5 ตาราง admin_user_categories — สิทธิ์หมวดของผู้จัดการ", size=15, bold=True, space_after=4)
    add_table(
        doc,
        ["ฟิลด์", "ชนิดข้อมูล", "คีย์/ข้อกำหนด", "คำอธิบาย"],
        [
            ["admin_user_id", "INT UNSIGNED", "PK, FK → admin_users.id", "ผู้จัดการ (ON DELETE CASCADE)"],
            ["category_id", "INT UNSIGNED", "PK, FK → categories.id", "หมวดที่ได้รับมอบ (ON DELETE CASCADE)"],
        ],
    )
    add_para(
        doc,
        "ความสัมพันธ์โดยสรุป: 1 หมวดมีได้หลาย FAQ · 1 FAQ มีเวกเตอร์ได้ 1 แถว · "
        "ผู้จัดการหลายคนดูแลได้หลายหมวดผ่านตารางเชื่อม admin_user_categories · "
        "การลบหมวดที่มี FAQ ค้างอยู่จะถูกกันไว้ (RESTRICT) เพื่อไม่ให้ข้อมูลคำตอบหายโดยไม่ตั้งใจ",
        size=15,
    )

    add_para(doc, "3.5.6 ข้อมูลนอกฐาน scigeniedb ที่ระบบอ้างถึง", size=15, bold=True, space_after=4)
    add_table(
        doc,
        ["แหล่งข้อมูล", "ใช้ทำอะไร", "หมายเหตุ"],
        [
            ["eoffice.t_user (อ่านอย่างเดียว)", "ตรวจว่าอีเมลเป็นบุคลากรที่ยังปฏิบัติงาน", "ไม่ใช่ FK จริงใน scigeniedb"],
            ["data/*.jsonl (AskLog)", "บันทึกประวัติการถาม–ตอบเพื่อสถิติ", "แยกจากตาราง FAQ"],
            ["ไฟล์ session ช่องทาง LINE/FB", "จำหมวดของผู้ใช้ระยะสั้น", "อยู่นอกเว็บที่ดาวน์โหลดตรงได้"],
        ],
    )

    add_h(doc, "3.6 หลักการออกแบบที่ห้ามทำลาย", 2)
    add_numbered(
        doc,
        [
            "แหล่งความจริงเดียว = คลัง FAQ ในฐานข้อมูล",
            "มีเกณฑ์ความเชื่อมั่น และกล้าไม่ตอบเมื่อหลักฐานไม่พอ",
            "ส่งต่อเจ้าหน้าที่เมื่อเกินขอบเขตคลัง",
            "สิทธิ์น้อยที่สุด (manager เห็นเฉพาะหมวดของตน)",
            "ลดการเก็บข้อมูลส่วนบุคคลเท่าที่จำเป็น",
        ],
    )

    # ===== Ch4 =====
    add_h(doc, "บทที่ 4  พัฒนาและเชื่อมต่อระบบ (Implementation)", 1)
    add_h(doc, "4.1 สแต็กเทคโนโลยี", 2)
    add_table(
        doc,
        ["ส่วน", "เทคโนโลยี"],
        [
            ["Backend", "PHP 8+"],
            ["Web server", "Apache (XAMPP / Docker)"],
            ["Database", "MySQL / MariaDB (utf8mb4)"],
            ["Auth ผู้ดูแล", "Google OAuth + ตรวจจากฐาน eoffice"],
            ["ช่องทาง", "LINE Messaging API, Facebook Messenger"],
            ["AI", "OpenAI-compatible gateway (ตรวจคำตอบ)"],
        ],
    )
    add_h(doc, "4.2 โครงสร้างโฟลเดอร์สำคัญ", 2)
    add_table(
        doc,
        ["เส้นทาง", "ใช้ทำอะไร"],
        [
            ["api/ask.php", "API ถาม–ตอบของเว็บ"],
            ["api/line-webhook.php", "รับข้อความ LINE"],
            ["api/fb-webhook.php", "รับข้อความ Facebook"],
            ["includes/Matcher.php", "ค้นหา + ให้คะแนน + ข้ามหมวด"],
            ["includes/OpenAiFaqVerifier.php", "AI ตรวจคำตอบ"],
            ["includes/LineBot.php / MessengerBot.php", "เชื่อมช่องทาง"],
            ["includes/Auth.php / FaqService.php", "สิทธิ์และคลัง FAQ"],
            ["admin/", "หน้าผู้ดูแล"],
            ["assets/js/widget.js", "วิดเจ็ตแชทบนเว็บ"],
            ["sql/schema.sql", "โครงสร้างฐานข้อมูล"],
            [".env", "ค่าความลับและค่าสภาพแวดล้อม"],
        ],
    )
    add_h(doc, "4.3 การรันบนเครื่องพัฒนา", 2)
    add_para(doc, "วิธีที่ 1 — Docker", bold=True, size=15)
    add_para(doc, "docker compose up --build -d   แล้วเปิด http://localhost:8080/", size=15)
    add_para(doc, "วิธีที่ 2 — XAMPP", bold=True, size=15)
    add_para(
        doc,
        "วางโปรเจกต์ใน htdocs เปิด Apache/MySQL แล้วเปิด http://localhost/SCiGENIE/ "
        "สร้างฐานด้วย sql/schema.sql และตั้งค่าใน .env",
        size=15,
    )
    add_h(doc, "4.4 ค่า .env ที่ต้องตั้ง (อย่า commit ความลับ)", 2)
    add_bullets(
        doc,
        [
            "ฐานข้อมูล: SCIGENIE_DB_* และ EOFFICE_DB_*",
            "Google OAuth: GOOGLE_CLIENT_ID / SECRET / REDIRECT_URI",
            "LINE: LINE_CHANNEL_SECRET / ACCESS_TOKEN / WEBHOOK_URL",
            "Facebook: FACEBOOK_PAGE_ACCESS_TOKEN / APP_SECRET / VERIFY_TOKEN",
            "AI: OPENAI_API_KEY / BASE_URL / MODEL (embedding ใช้ local-hash-v2 ได้)",
        ],
    )
    add_h(doc, "4.5 กฎการเขียนโค้ดสำหรับผู้รับช่วง", 2)
    add_bullets(
        doc,
        [
            "ทุกช่องทางต้องเรียก Matcher/FaqService ตัวเดียวกัน",
            "อย่าฝัง API key หรือ token ในโค้ด",
            "คำสั่ง SQL ใช้ prepared statement เสมอ",
            "ตรวจสิทธิ์หมวดฝั่งเซิร์ฟเวอร์ทุกครั้งที่เพิ่ม/แก้/ลบ FAQ",
            "หลังแก้ FAQ จำนวนมาก ให้รัน tools/rebuild-embeddings.php ตามความจำเป็น",
        ],
    )
    add_h(doc, "4.6 เชื่อมต่อช่องทางภายนอก (สรุป)", 2)
    add_para(doc, "LINE", bold=True, size=15)
    add_bullets(
        doc,
        [
            "Webhook: .../api/line-webhook.php",
            "ตรวจลายเซ็นด้วย Channel Secret ทุกครั้ง",
            "ตอบรับเร็ว แล้วค่อยประมวลผลคำถาม",
        ],
    )
    add_para(doc, "Facebook Messenger", bold=True, size=15)
    add_bullets(
        doc,
        [
            "Webhook: .../api/fb-webhook.php",
            "Verify Token ต้องตรงกับ FACEBOOK_VERIFY_TOKEN",
            "ตรวจลายเซ็นด้วย App Secret ก่อนประมวลผล",
            "โหมดพัฒนา: ตอบได้เฉพาะ Admin/Tester จนกว่าแอปจะ Live",
        ],
    )
    add_para(doc, "Google Login (เฉพาะแอดมิน)", bold=True, size=15)
    add_bullets(
        doc,
        [
            "ลำดับตรวจ: Google email → พบใน eoffice และยังปฏิบัติงาน → มีใน admin_users และเปิดใช้",
            "ผู้รับบริการถามผ่านเว็บ/LINE/FB ไม่ต้องล็อกอิน",
        ],
    )

    # ===== Ch5 =====
    add_h(doc, "บทที่ 5  ทดสอบระบบ (Testing)", 1)
    add_h(doc, "5.1 สิ่งที่ต้องทดสอบเป็นพิเศษ", 2)
    add_para(
        doc,
        "อย่าหยุดแค่ “หน้าจอเปิดได้” ความเสียหายหลักคือ “ตอบผิดอย่างมั่นใจ” หรือ “ตอบคนละหมวดระหว่างช่องทาง”",
    )
    add_h(doc, "5.2 ชุดทดสอบขั้นต่ำ", 2)
    add_table(
        doc,
        ["ชุด", "ตัวอย่างเคส", "ผลที่คาดหวัง"],
        [
            ["ภาษาทางการ", "คำถามใกล้หัวข้อในคลัง", "ตอบถูกหมวด"],
            ["คำสั้น / คำย่อ", "พสวท. / เบอร์โทรคณะ", "จับคู่ได้ หรือขอหมวด"],
            ["ข้ามหมวด", "เลือกหมวด A แล้วถามเรื่องหมวด B", "ค้นข้ามหมวดแล้วตอบถูก"],
            ["ไม่มีในคลัง", "เรื่องที่ตั้งใจไม่มี FAQ", "ไม่เดา + แนะนำเจ้าหน้าที่"],
            ["สิทธิ์แอดมิน", "manager แก้หมวดอื่น", "ถูกปฏิเสธ"],
            ["ช่องทาง", "ถามชุดเดียวกันบนเว็บ/LINE/FB", "เนื้อหาคำตอบชุดเดียวกัน"],
        ],
    )
    add_h(doc, "5.3 ถดถอย (Regression)", 2)
    add_para(
        doc,
        "ทุกครั้งที่เพิ่มคำพ้อง เปลี่ยนน้ำหนักคะแนน หรือแก้ FAQ จำนวนมาก ให้รันชุดเดิมซ้ำอย่างน้อยในหมวดที่กระทบ",
    )

    # ===== Ch6 =====
    add_h(doc, "บทที่ 6  ติดตั้งขึ้นเครื่องแม่ข่าย (Deployment)", 1)
    add_numbered(
        doc,
        [
            "สร้างฐาน scigeniedb จาก sql/schema.sql (และ faq_embeddings ถ้ายังไม่มี)",
            "วางโค้ดบนเซิร์ฟเวอร์ ตั้งสิทธิ์โฟลเดอร์ data/ ให้เขียนได้",
            "ตั้งค่า .env ให้ครบ (DB, OAuth, LINE/FB, AI)",
            "กำหนดผู้ดูแลระบบคนแรก แล้วให้ล็อกอินด้วย Google",
            "ตั้ง webhook LINE/Facebook ให้ชี้มาที่โดเมนจริง (HTTPS)",
            "ตรวจ privacy.php เปิดได้สาธารณะ",
            "รันชุดทดสอบบน production อย่างน้อย 1 รอบก่อนเปิดใช้กว้าง",
        ],
    )
    add_para(
        doc,
        "URL ใช้งานจริง: https://sci.kku.ac.th/app/SCiGENIE/   |   แอดมิน: .../admin/",
        size=15,
    )

    # ===== Ch7 =====
    add_h(doc, "บทที่ 7  บำรุงรักษาและปรับปรุงต่อเนื่อง (Maintenance)", 1)
    add_h(doc, "7.1 งานประจำของเจ้าของหมวด", 2)
    add_bullets(
        doc,
        [
            "ดูสถิติคำถามที่ตอบไม่ได้ แล้วเติม FAQ",
            "อัปเดตคำตอบเมื่อระเบียบ/ประกาศเปลี่ยน",
            "เติมคำสำคัญเสริมจากภาษาพูดที่ผู้ใช้ถามจริง",
        ],
    )
    add_h(doc, "7.2 งานประจำของผู้พัฒนา/ผู้ดูแลระบบ", 2)
    add_bullets(
        doc,
        [
            "ตรวจเว็บเซิร์ฟเวอร์ ฐานข้อมูล และ error log ของ webhook",
            "สำรองฐานข้อมูลและไฟล์ log เป็นระยะ",
            "หมุนเวียน secret/token เมื่อมีการเปลี่ยนผู้ดูแลเครื่อง",
            "หลังเปลี่ยนรุ่น embedding หรือโมเดล AI ให้ rebuild ดัชนีและทดสอบความแม่นยำ",
        ],
    )
    add_h(doc, "7.3 ข้อจำกัดที่ตั้งใจไว้", 2)
    add_bullets(
        doc,
        [
            "ตอบได้เฉพาะที่มีในคลัง",
            "AI เป็นผู้ตรวจ ไม่ใช่ผู้แต่งคำตอบ",
            "สถานะหมวดบน LINE/FB เป็นสถานะระยะสั้นรายผู้ใช้",
        ],
    )
    add_h(doc, "7.4 สิ่งที่ไม่ควรทำในระยะใกล้", 2)
    add_bullets(
        doc,
        [
            "อย่าให้โมเดลภาษาตอบอิสระโดยไม่มีรายการจากคลัง",
            "อย่าลดเกณฑ์คะแนนเพียงเพื่อให้ตอบมากขึ้น โดยไม่มีชุดทดสอบกำกับ",
            "อย่าใช้ค่าลับของ production บนเครื่องพัฒนาสาธารณะ",
        ],
    )

    # ===== Appendix =====
    add_h(doc, "ภาคผนวก  ตารางอ้างอิงด่วน", 1)
    add_h(doc, "ก. บทบาทผู้ใช้ระบบหลังบ้าน", 2)
    add_table(
        doc,
        ["บทบาท", "สิทธิ์"],
        [
            ["super", "ทุกหมวด, จัดการผู้ใช้, ดูสถิติรวม, ดูประวัติแอดมิน"],
            ["manager", "FAQ + สถิติเฉพาะหมวดที่ได้รับมอบ"],
        ],
    )
    add_h(doc, "ข. Checklist ก่อนขึ้น production", 2)
    add_bullets(
        doc,
        [
            "☐ .env ครบและไม่อยู่ใน git",
            "☐ ฐานข้อมูล + embeddings พร้อม",
            "☐ Google OAuth redirect ตรงโดเมนจริง",
            "☐ LINE/FB webhook ตรวจลายเซ็นผ่าน",
            "☐ ผู้ดูแลคนแรกเข้าได้",
            "☐ ทดสอบถามถูก/ไม่ถูก/ข้ามหมวด",
            "☐ privacy.php เปิดได้",
        ],
    )
    add_h(doc, "ค. ไฟล์เครื่องมือช่วย", 2)
    add_bullets(
        doc,
        [
            "tools/rebuild-embeddings.php — สร้างดัชนีความหมายใหม่",
            "tools/test-match.php / test-dpst-match.php — ทดสอบการจับคู่",
            "sql/schema.sql / sql/faq_embeddings.sql — สคีมาฐานข้อมูล",
        ],
    )

    add_para(doc, "", space_after=12)
    add_para(
        doc,
        "สรุป: พัฒนาตาม SDLC นี้โดยคงหลัก “คลังเดียว เกณฑ์คะแนนชัด AI ช่วยตรวจ ไม่เดาเมื่อไม่มั่นใจ” "
        "จะทำให้ระบบดูแลต่อได้และคำตอบยังตรงกับข้อมูลของคณะ",
        size=15,
        bold=True,
    )

    os.makedirs("HR", exist_ok=True)
    doc.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    main()
